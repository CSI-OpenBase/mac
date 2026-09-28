from __future__ import annotations

import json
import re
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import admin_app.local_app as local_app_module
import admin_app.local_store as local_store_module
from admin_app import __version__
from admin_app.local_app import create_local_app
from admin_app.local_cleanup import ClearDataResult
from admin_app.local_config import (
    COMMENT_EXPORT_DIRECTORY_KEY,
    LOCAL_PREFERENCES_META_KEY,
    LocalSettings,
)
from admin_app.local_store import LocalStore


class FakeRunner:
    def __init__(self, store: LocalStore) -> None:
        self.store = store
        self.calls: list[tuple[str, dict[str, Any]]] = []

    def start(self) -> None:
        return None

    def submit(self, kind: str, **payload: Any) -> dict[str, Any]:
        self.calls.append((kind, payload))
        return self.store.create_job(
            kind, video_id=payload.get("video_id"), payload=payload
        )


def settings(
    tmp_path: Path, token: str = "", instance_nonce: str = ""
) -> LocalSettings:
    return LocalSettings(
        data_home=tmp_path,
        session_home=tmp_path / ".sessions",
        desktop_token=token,
        instance_nonce=instance_nonce,
    )


def csrf(client: TestClient) -> str:
    response = client.get("/")
    marker = 'name="csrf_token" value="'
    return response.text.split(marker, 1)[1].split('"', 1)[0]


def archive_videos(store: LocalStore, count: int) -> list[str]:
    records: list[dict[str, object]] = []
    video_ids: list[str] = []
    for index in range(count):
        video_id = str(7_700_000_000_000_000_000 + index)
        video_ids.append(video_id)
        records.append(
            {
                "video_id": video_id,
                "title": f"分页视频 {index:03d}",
                "video_url": f"https://www.douyin.com/video/{video_id}",
                "manifest_path": f"works/videos/douyin/{video_id}/manifest.json",
                "first_seen_at": "2026-09-07T10:00:00Z",
                "last_seen_at": (
                    f"2026-09-07T{10 + index // 60:02d}:{index % 60:02d}:00Z"
                ),
            }
        )
    store.upsert_videos(records)
    return video_ids


def rendered_video_ids(html: str) -> list[str]:
    return re.findall(r'name="video_id" value="(\d+)"', html)


def test_local_home_and_manual_authorization_job(tmp_path: Path) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    runner = FakeRunner(store)
    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        response = client.get("/")
        assert response.status_code == 200
        assert "本地归档" in response.text
        assert "brand/openbase-mark-reversed.svg" in response.text
        favicon_response = client.get("/favicon.ico")
        assert favicon_response.status_code == 200
        assert favicon_response.headers["content-type"] == "image/x-icon"
        assert favicon_response.content.startswith(b"\x00\x00\x01\x00")
        token = csrf(client)
        response = client.post(
            "/actions/authorize", data={"csrf_token": token}, follow_redirects=False
        )
        assert response.status_code == 303
    assert runner.calls == [("authorize", {})]


def test_local_home_displays_stored_utc_times_in_beijing(
    tmp_path: Path, monkeypatch
) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    monkeypatch.setattr(
        local_store_module, "utc_now", lambda: "2026-09-07T04:00:00Z"
    )
    store.create_job("authorize")
    archive_videos(store, 1)
    store.set_meta(
        "last_export",
        {"finished_at": "2026-09-07T04:00:00Z", "summary": "完成"},
    )

    with TestClient(
        create_local_app(local_settings, store=store, runner=FakeRunner(store))
    ) as client:
        home = client.get("/")
        tasks = client.get("/tasks")

    assert home.status_code == 200
    assert tasks.status_code == 200
    assert 'href="/tasks"' in home.text
    assert "开始时间（北京时间）" not in home.text
    assert "2026-09-07 12:00:00" in home.text
    assert "2026-09-07 18:00:00" in home.text
    assert "开始时间（北京时间）" in tasks.text
    assert "2026-09-07 12:00:00" in tasks.text


def test_local_home_displays_indexed_and_legacy_video_publish_times(
    tmp_path: Path,
) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    indexed_video_id, legacy_video_id = archive_videos(store, 2)
    store.upsert_videos(
        [
            {
                "video_id": indexed_video_id,
                "title": "索引发布时间",
                "video_url": f"https://www.douyin.com/video/{indexed_video_id}",
                "manifest_path": (
                    f"works/videos/douyin/{indexed_video_id}/manifest.json"
                ),
                "first_seen_at": "2026-09-07T10:00:00Z",
                "last_seen_at": "2026-09-07T10:00:00Z",
                "published_at": "2026-09-01T00:30:00Z",
            }
        ]
    )
    legacy_directory = local_settings.videos_dir / "douyin" / legacy_video_id
    metadata_path = legacy_directory / "metadata" / "latest.json"
    metadata_path.parent.mkdir(parents=True)
    metadata_path.write_text(
        json.dumps({"published_at": "2026-09-02T01:00:00Z"}),
        encoding="utf-8",
    )
    (legacy_directory / "manifest.json").write_text(
        json.dumps({"latest_metadata": "metadata/latest.json"}),
        encoding="utf-8",
    )

    with TestClient(
        create_local_app(local_settings, store=store, runner=FakeRunner(store))
    ) as client:
        response = client.get("/")

    assert response.status_code == 200
    assert "<th>发布时间</th>" in response.text
    indexed_row = response.text.split(
        f'id="video-{indexed_video_id}"', 1
    )[1].split("</tr>", 1)[0]
    legacy_row = response.text.split(
        f'id="video-{legacy_video_id}"', 1
    )[1].split("</tr>", 1)[0]
    assert 'datetime="2026-09-01T00:30:00Z"' in indexed_row
    assert "2026-09-01 08:30:00" in indexed_row
    assert 'datetime="2026-09-02T01:00:00Z"' in legacy_row
    assert "2026-09-02 09:00:00" in legacy_row


def test_video_archive_page_lists_synced_metadata_and_serves_safe_cover(
    tmp_path: Path,
) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    video_ids = archive_videos(store, 31)
    video_id = video_ids[-1]
    video_directory = local_settings.videos_dir / "douyin" / video_id
    metadata_path = video_directory / "metadata" / "latest.json"
    metadata_path.parent.mkdir(parents=True)
    metadata_path.write_text(
        json.dumps(
            {
                "desc": "完整作品描述",
                "published_at": "2026-09-01T00:30:00Z",
                "observed_at": "2026-09-07T04:00:00Z",
                "visible_metrics": {
                    "view_count": 12_345,
                    "like_count": 678,
                    "comment_count": 90,
                    "collect_count": 45,
                    "share_count": 12,
                },
                "sources": ["response", "dom"],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    (video_directory / "manifest.json").write_text(
        json.dumps({"latest_metadata": "metadata/latest.json"}),
        encoding="utf-8",
    )
    cover_path = video_directory / "cover.jpg"
    cover_path.write_bytes(b"\xff\xd8\xff\xd9")
    store.upsert_videos(
        [
            {
                "video_id": video_id,
                "title": "完整档案作品",
                "video_url": f"https://www.douyin.com/video/{video_id}",
                "cover_path": cover_path.relative_to(
                    local_settings.data_home
                ).as_posix(),
                "manifest_path": (
                    f"works/videos/douyin/{video_id}/manifest.json"
                ),
                "first_seen_at": "2026-09-07T01:00:00Z",
                "last_seen_at": "2026-09-07T02:00:00Z",
                "platform_groups_observed": True,
                "platform_groups": [{"id": "column-1", "name": "档案栏目"}],
            }
        ]
    )

    with TestClient(
        create_local_app(local_settings, store=store, runner=FakeRunner(store))
    ) as client:
        home = client.get("/")
        archive = client.get("/video-archive?page_size=30")
        second_page = client.get("/video-archive?page=2&page_size=30")
        cover = client.get(f"/video-covers/{video_id}")
        invalid_cover = client.get("/video-covers/not-a-video")

    assert home.status_code == 200
    assert 'href="/video-archive"' in home.text
    assert "查看视频档案" in home.text
    assert archive.status_code == 200
    assert "31 个已同步作品" in archive.text
    assert 'action="/video-archive/export"' in archive.text
    assert archive.text.count("必选") == 3
    assert "<th>视频</th><th>标题</th><th>描述</th>" in archive.text
    assert 'class="video-catalog-title"' in archive.text
    assert "完整档案作品" in archive.text
    assert '<td class="video-catalog-description">完整作品描述</td>' in archive.text
    assert "2026-09-01 08:30:00" in archive.text
    assert "档案栏目" in archive.text
    assert "12,345" in archive.text
    assert "678" in archive.text
    assert "90" in archive.text
    assert "45" in archive.text
    assert "平台接口 / 页面可见内容" in archive.text
    assert "观测 2026-09-07 12:00:00" in archive.text
    assert "/video-archive?page_size=30&amp;page=2" in archive.text
    assert second_page.status_code == 200
    assert video_ids[0] in second_page.text
    assert cover.status_code == 200
    assert cover.headers["content-type"] == "image/jpeg"
    assert cover.content == b"\xff\xd8\xff\xd9"
    assert invalid_cover.status_code == 404


def test_video_archive_export_enforces_required_columns_and_selected_fields(
    tmp_path: Path,
) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    video_id = "7700000000000000001"
    video_directory = local_settings.videos_dir / "douyin" / video_id
    metadata_path = video_directory / "metadata" / "latest.json"
    metadata_path.parent.mkdir(parents=True)
    metadata_path.write_text(
        json.dumps(
            {
                "desc": "单独导出的描述",
                "published_at": "2026-09-01T00:30:00Z",
                "observed_at": "2026-09-07T04:00:00Z",
                "visible_metrics": {"view_count": 12_345},
                "sources": ["response"],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    (video_directory / "manifest.json").write_text(
        json.dumps({"latest_metadata": "metadata/latest.json"}),
        encoding="utf-8",
    )
    store.upsert_videos(
        [
            {
                "video_id": video_id,
                "title": "=作为文本的标题",
                "video_url": f"https://www.douyin.com/video/{video_id}",
                "manifest_path": (
                    f"works/videos/douyin/{video_id}/manifest.json"
                ),
                "first_seen_at": "2026-09-07T01:00:00Z",
                "last_seen_at": "2026-09-07T02:00:00Z",
                "platform_groups_observed": True,
                "platform_groups": [{"id": "column-1", "name": "档案栏目"}],
            }
        ]
    )

    with TestClient(
        create_local_app(local_settings, store=store, runner=FakeRunner(store))
    ) as client:
        token = csrf(client)
        exported = client.post(
            "/video-archive/export",
            data={
                "csrf_token": token,
                "columns": ["description", "groups", "view_count"],
            },
        )
        invalid = client.post(
            "/video-archive/export",
            data={"csrf_token": token, "columns": ["author"]},
        )
        missing_csrf = client.post(
            "/video-archive/export",
            data={"columns": ["description"]},
        )

    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert re.fullmatch(
        r'attachment; filename="csi-openbase-video-archive-\d{4}-\d{2}-\d{2}_\d{2}-\d{2}-\d{2}\.xlsx"',
        exported.headers["content-disposition"],
    )
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    rows = list(workbook["视频档案"].iter_rows(values_only=False))
    assert [cell.value for cell in rows[0]] == [
        "视频 ID",
        "标题",
        "发布时间",
        "描述",
        "分组",
        "播放数",
    ]
    assert [cell.value for cell in rows[1]] == [
        video_id,
        "=作为文本的标题",
        "2026-09-01 08:30:00",
        "单独导出的描述",
        "档案栏目",
        12_345,
    ]
    assert rows[1][0].data_type == "s"
    assert rows[1][1].data_type == "s"
    assert invalid.status_code == 400
    assert missing_csrf.status_code == 403


def test_local_home_marks_comment_export_status_per_video(tmp_path: Path) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    exported_video_id, pending_video_id = archive_videos(store, 2)
    store.record_comment_export(
        exported_video_id,
        count=27,
        exported_at="2026-09-07T12:00:00Z",
    )

    with TestClient(
        create_local_app(local_settings, store=store, runner=FakeRunner(store))
    ) as client:
        response = client.get("/")

    assert response.status_code == 200
    exported_row = response.text.split(
        f'id="video-{exported_video_id}"', 1
    )[1].split("</tr>", 1)[0]
    pending_row = response.text.split(
        f'id="video-{pending_video_id}"', 1
    )[1].split("</tr>", 1)[0]
    assert 'data-comment-export-status="exported"' in exported_row
    assert 'data-lucide="circle-check"' in exported_row
    assert "已导出" in exported_row
    assert "27 条 · 2026-09-07 20:00:00" in exported_row
    assert 'data-comment-export-status="not-exported"' in pending_row
    assert 'data-lucide="circle-dashed"' in pending_row
    assert "未导出" in pending_row


def test_task_history_has_a_dedicated_page(tmp_path: Path) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    job_id = store.create_job("authorize")["id"]
    video_id = archive_videos(store, 1)[0]
    video_job_id = store.create_job("comments", video_id=video_id)["id"]

    with TestClient(
        create_local_app(local_settings, store=store, runner=FakeRunner(store))
    ) as client:
        home = client.get("/")
        tasks = client.get("/tasks")
        settings_page = client.get("/settings")

    assert home.status_code == 200
    assert tasks.status_code == 200
    assert settings_page.status_code == 200
    assert 'href="/tasks"' in home.text
    assert 'href="/tasks"' in settings_page.text
    assert "最近 30 条由用户主动创建的采集任务" in tasks.text
    assert f"#{job_id}" in tasks.text
    assert f"#{video_job_id}" in tasks.text
    assert f'href="/videos/{video_id}"' in tasks.text
    assert "暂无任务记录" not in tasks.text
    assert "开始时间（北京时间）" in tasks.text
    assert "history-band" not in home.text


def test_video_task_link_redirects_to_the_archived_video_page(tmp_path: Path) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    video_ids = archive_videos(store, 105)

    with TestClient(
        create_local_app(local_settings, store=store, runner=FakeRunner(store))
    ) as client:
        redirect = client.get(f"/videos/{video_ids[0]}", follow_redirects=False)
        target = client.get(redirect.headers["location"])
        missing = client.get("/videos/12345678", follow_redirects=False)
        invalid = client.get("/videos/not-a-video", follow_redirects=False)

    assert redirect.status_code == 303
    assert redirect.headers["location"] == (
        f"/?page=4&page_size=30#video-{video_ids[0]}"
    )
    assert target.status_code == 200
    assert f'id="video-{video_ids[0]}"' in target.text
    assert missing.status_code == 303
    assert missing.headers["location"] == "/#video-archive"
    assert invalid.status_code == 404


def test_comment_export_directory_setting_can_be_saved_and_reset(
    tmp_path: Path,
) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    runner = FakeRunner(store)
    selected = tmp_path / "user-comment-exports"
    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        home = client.get("/")
        assert 'href="/settings"' in home.text
        page = client.get("/settings")
        assert page.status_code == 200
        assert "评论导出目录" in page.text
        assert "data-comment-directory-picker" in page.text

        response = client.post(
            "/settings/comments",
            data={
                "csrf_token": csrf(client),
                "comment_export_directory": str(selected),
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert response.headers["location"] == "/settings"
        assert store.get_meta(LOCAL_PREFERENCES_META_KEY, {}) == {
            COMMENT_EXPORT_DIRECTORY_KEY: str(selected.resolve())
        }
        assert str(selected.resolve()) in client.get("/settings").text

        response = client.post(
            "/settings/comments",
            data={
                "csrf_token": csrf(client),
                "comment_export_directory": "",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert store.get_meta(LOCAL_PREFERENCES_META_KEY, {}) == {
            COMMENT_EXPORT_DIRECTORY_KEY: ""
        }


def test_local_home_paginates_video_archive_with_selectable_page_size(
    tmp_path: Path,
) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    video_ids = archive_videos(store, 105)
    runner = FakeRunner(store)

    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        first = client.get("/")
        second = client.get("/?page=2&page_size=50")
        last = client.get("/?page=99&page_size=50")
        hundred = client.get("/?page=2&page_size=100")
        invalid = client.get("/?page=invalid&page_size=25")

    assert first.status_code == 200
    assert rendered_video_ids(first.text) == list(reversed(video_ids[75:]))
    assert "105 个视频" in first.text
    assert '<option value="30" selected>30</option>' in first.text
    assert 'aria-label="选择本页全部视频"' in first.text
    assert 'aria-label="视频档案分页"' in first.text
    assert '/?page_size=30&amp;page=2#video-archive' in first.text

    assert rendered_video_ids(second.text) == list(reversed(video_ids[5:55]))
    assert '<option value="50" selected>50</option>' in second.text
    assert '<span class="page-number is-current" aria-current="page">2</span>' in second.text

    assert rendered_video_ids(last.text) == list(reversed(video_ids[:5]))
    assert '<span class="page-number is-current" aria-current="page">3</span>' in last.text
    assert 'class="page-button is-disabled" aria-disabled="true"><span>下一页' in last.text

    assert rendered_video_ids(hundred.text) == list(reversed(video_ids[:5]))
    assert '<option value="100" selected>100</option>' in hundred.text
    assert rendered_video_ids(invalid.text) == list(reversed(video_ids[75:]))


def test_local_home_creates_filters_and_manages_manual_video_groups(
    tmp_path: Path,
) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    video_ids = archive_videos(store, 2)
    runner = FakeRunner(store)

    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        token = csrf(client)
        created = client.post(
            "/groups",
            data={"csrf_token": token, "name": "待复盘"},
            follow_redirects=False,
        )
        group = store.list_groups()[0]
        assert created.status_code == 303
        assert group["name"] == "待复盘"

        added = client.post(
            "/groups/videos/add",
            data={
                "csrf_token": token,
                "target_group_id": group["group_id"],
                "video_id": [video_ids[0]],
                "page": "1",
                "page_size": "30",
            },
            follow_redirects=False,
        )
        assert added.status_code == 303
        filtered = client.get(
            "/",
            params={"group_id": group["group_id"], "page_size": 30},
        )
        assert rendered_video_ids(filtered.text) == [video_ids[0]]
        assert ">待复盘</a>" in filtered.text
        assert ">移出分组</span>" in filtered.text
        assert 'formaction="/groups/videos/add"' in filtered.text

        renamed = client.post(
            f"/groups/{group['group_id']}/rename",
            data={
                "csrf_token": token,
                "name": "重点作品",
                "group_id": group["group_id"],
                "page": "1",
                "page_size": "30",
            },
            follow_redirects=False,
        )
        assert renamed.status_code == 303
        assert store.get_group(group["group_id"])["name"] == "重点作品"

        removed = client.post(
            "/groups/videos/remove",
            data={
                "csrf_token": token,
                "group_id": group["group_id"],
                "video_id": [video_ids[0]],
                "page": "1",
                "page_size": "30",
            },
            follow_redirects=False,
        )
        assert removed.status_code == 303
        assert store.get_group(group["group_id"])["video_count"] == 0

        deleted = client.post(
            f"/groups/{group['group_id']}/delete",
            data={"csrf_token": token},
            follow_redirects=False,
        )
        assert deleted.status_code == 303
        assert store.list_groups() == []


def test_local_home_exposes_scoped_clear_dialog(tmp_path: Path) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    runner = FakeRunner(store)

    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        response = client.get("/")

    assert response.status_code == 200
    assert client.app.version == __version__
    assert f'<span class="brand-version">v{__version__}</span>' in response.text
    assert ">清空数据</span>" in response.text
    assert 'action="/actions/clear-data"' in response.text
    assert 'name="scope" value="exports" checked' in response.text
    assert 'name="scope" value="comments"' in response.text
    assert 'name="scope" value="all"' in response.text
    assert "保留登录授权" in response.text


def test_clear_data_route_requires_confirmation_and_clears_selected_scope(
    tmp_path: Path,
) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    store.set_meta("last_export", {"directory": "exports/run-1"})
    export_file = local_settings.exports_dir / "run-1" / "table.xlsx"
    export_file.parent.mkdir(parents=True)
    export_file.write_bytes(b"export")
    runner = FakeRunner(store)

    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        token = csrf(client)
        rejected = client.post(
            "/actions/clear-data",
            data={"csrf_token": token, "scope": "exports"},
            follow_redirects=False,
        )
        assert rejected.status_code == 303
        assert export_file.is_file()
        assert "请先确认清空操作无法撤销" in client.get("/").text

        accepted = client.post(
            "/actions/clear-data",
            data={
                "csrf_token": csrf(client),
                "scope": "exports",
                "confirm_clear": "yes",
            },
            follow_redirects=False,
        )
        assert accepted.status_code == 303
        home = client.get("/")

    assert not export_file.exists()
    assert local_settings.exports_dir.is_dir()
    assert store.get_meta("last_export", {}) == {}
    assert "已清空平台导出的原始数据" in home.text


def test_clear_data_route_requires_csrf(tmp_path: Path) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    export_file = local_settings.exports_dir / "run-1" / "table.xlsx"
    export_file.parent.mkdir(parents=True)
    export_file.write_bytes(b"export")
    runner = FakeRunner(store)

    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        response = client.post(
            "/actions/clear-data",
            data={"scope": "exports", "confirm_clear": "yes"},
        )

    assert response.status_code == 403
    assert export_file.is_file()


def test_clear_data_route_warns_when_physical_deletion_is_pending(
    tmp_path: Path, monkeypatch
) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    runner = FakeRunner(store)
    monkeypatch.setattr(
        local_app_module,
        "clear_local_data",
        lambda *_args: ClearDataResult(
            scope="exports",
            files_deleted=1,
            directories_deleted=1,
            pending_directories=1,
        ),
    )

    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        response = client.post(
            "/actions/clear-data",
            data={
                "csrf_token": csrf(client),
                "scope": "exports",
                "confirm_clear": "yes",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        home = client.get("/")

    assert "暂存目录因文件占用未能删除" in home.text


def test_clear_data_route_rejects_active_jobs(tmp_path: Path) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    export_file = local_settings.exports_dir / "run-1" / "table.xlsx"
    export_file.parent.mkdir(parents=True)
    export_file.write_bytes(b"export")
    runner = FakeRunner(store)

    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        store.create_job("authorize")
        response = client.post(
            "/actions/clear-data",
            data={
                "csrf_token": csrf(client),
                "scope": "exports",
                "confirm_clear": "yes",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        home = client.get("/")

    assert export_file.is_file()
    assert "当前仍有任务等待或运行" in home.text


def test_export_requires_authorization(tmp_path: Path) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    runner = FakeRunner(store)
    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        response = client.post(
            "/actions/export",
            data={"csrf_token": csrf(client)},
            follow_redirects=False,
        )
        assert response.status_code == 303
    assert runner.calls == []


def test_complete_export_requires_a_profile_sync(tmp_path: Path) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    store.set_meta("creator_identity", {"handle": "creator-handle"})
    runner = FakeRunner(store)
    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        response = client.post(
            "/actions/export",
            data={"csrf_token": csrf(client)},
            follow_redirects=False,
        )
        assert response.status_code == 303
    assert runner.calls == []


def test_complete_export_rejects_an_incomplete_profile_archive(
    tmp_path: Path,
) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    store.set_meta("creator_identity", {"handle": "creator-handle"})
    store.set_meta(
        "last_video_sync",
        {"finished_at": "2026-09-07T12:00:00Z", "complete": False},
    )
    runner = FakeRunner(store)
    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        response = client.post(
            "/actions/export",
            data={"csrf_token": csrf(client)},
            follow_redirects=False,
        )
        assert response.status_code == 303
    assert runner.calls == []


def test_sync_route_always_targets_the_signed_in_profile(tmp_path: Path) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    store.set_meta("creator_identity", {"handle": "creator-handle"})
    runner = FakeRunner(store)
    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        response = client.post(
            "/actions/sync-videos",
            data={
                "csrf_token": csrf(client),
                "profile_url": "https://www.douyin.com/user/someone-else",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
    assert runner.calls == [("sync_videos", {})]


def test_desktop_token_guards_ui_and_identifies_health_instance(tmp_path: Path) -> None:
    local_settings = settings(
        tmp_path, token="desktop-secret", instance_nonce="instance-123"
    )
    store = LocalStore(local_settings.database_path)
    runner = FakeRunner(store)
    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        assert client.get("/health").status_code == 401
        assert client.get("/").status_code == 401
        assert client.get("/favicon.ico").status_code == 200
        response = client.get(
            "/health", headers={"X-CSI-Desktop-Token": "desktop-secret"}
        )
        assert response.status_code == 200
        assert response.json()["instance_nonce"] == "instance-123"
        assert response.json()["version"] == __version__
        assert client.get("/").status_code == 200


def test_desktop_token_replaces_a_stale_cookie_after_restart(tmp_path: Path) -> None:
    local_settings = settings(tmp_path, token="new-desktop-secret")
    store = LocalStore(local_settings.database_path)
    runner = FakeRunner(store)
    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        client.cookies.set(
            "csi_desktop",
            "old-desktop-secret",
            domain="testserver.local",
            path="/",
        )
        response = client.get(
            "/", headers={"X-CSI-Desktop-Token": "new-desktop-secret"}
        )
        assert response.status_code == 200
        assert client.cookies.get(
            "csi_desktop", domain="testserver.local", path="/"
        ) == "new-desktop-secret"
        assert client.get("/static/local.css").status_code == 200


def test_shutdown_requires_desktop_token_and_invokes_callback(tmp_path: Path) -> None:
    stopped: list[bool] = []
    local_settings = settings(tmp_path, token="desktop-secret")
    store = LocalStore(local_settings.database_path)
    runner = FakeRunner(store)
    with TestClient(
        create_local_app(
            local_settings,
            store=store,
            runner=runner,
            shutdown_callback=lambda: stopped.append(True),
        )
    ) as client:
        assert client.post("/api/shutdown").status_code == 401
        response = client.post(
            "/api/shutdown",
            headers={"X-CSI-Desktop-Token": "desktop-secret"},
        )
        assert response.status_code == 200
        assert response.json() == {
            "status": "stopping",
            "force_required": False,
        }
    assert stopped == [True]


def test_shutdown_is_not_exposed_in_source_mode(tmp_path: Path) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    runner = FakeRunner(store)
    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        assert client.post("/api/shutdown").status_code == 404


def test_shutdown_waits_for_desktop_to_force_cleanup_of_active_browser_job(
    tmp_path: Path, monkeypatch
) -> None:
    stopped: list[bool] = []

    class ActiveRunner:
        def __init__(self, *_args: Any, **_kwargs: Any) -> None:
            self.close_calls = 0

        def start(self) -> None:
            return None

        def close(self) -> bool:
            self.close_calls += 1
            return False

    import admin_app.local_jobs as jobs_module

    monkeypatch.setattr(jobs_module, "LocalJobRunner", ActiveRunner)
    local_settings = settings(
        tmp_path, token="desktop-secret", instance_nonce="instance-123"
    )
    with TestClient(
        create_local_app(
            local_settings,
            shutdown_callback=lambda: stopped.append(True),
        )
    ) as client:
        response = client.post(
            "/api/shutdown",
            headers={"X-CSI-Desktop-Token": "desktop-secret"},
        )
        assert response.status_code == 200
        assert response.json() == {
            "status": "stopping",
            "force_required": True,
        }
        assert stopped == []


def test_comment_route_only_creates_explicit_immediate_job(tmp_path: Path) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    store.upsert_videos(
        [
            {
                "video_id": "7680023068660346011",
                "title": "视频",
                "video_url": "https://www.douyin.com/video/7680023068660346011",
                "manifest_path": "works/videos/douyin/7680023068660346011/manifest.json",
                "first_seen_at": "2026-09-07T00:00:00Z",
                "last_seen_at": "2026-09-07T00:00:00Z",
                "visible_comment_count": 12_876,
            }
        ]
    )
    runner = FakeRunner(store)
    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        home = client.get("/")
        assert home.status_code == 200
        assert "12,876" in home.text
        assert "首次获取" in home.text
        assert "尚未导出评论" in home.text
        assert ">获取最新评论</span>" in home.text
        assert ">导出评论</span>" in home.text
        assert ">导出新增</span>" not in home.text
        assert "完整重新同步评论" in home.text
        assert (
            'formaction="/videos/7680023068660346011/comment-count"' in home.text
        )
        assert (
            'formaction="/videos/7680023068660346011/comments"' in home.text
        )
        assert (
            'formaction="/videos/7680023068660346011/comments/full"' in home.text
        )
        response = client.post(
            "/videos/7680023068660346011/comments",
            data={"csrf_token": csrf(client), "page": "2", "page_size": "50"},
            follow_redirects=False,
        )
        assert response.status_code == 303
        assert response.headers["location"] == (
            "/?page=2&page_size=50#video-archive"
        )
        first_job = store.list_jobs(limit=1)[0]
        store.update_job(first_job["id"], "succeeded")
        response = client.post(
            "/videos/7680023068660346011/comments/full",
            data={"csrf_token": csrf(client), "page": "2", "page_size": "50"},
            follow_redirects=False,
        )
        assert response.status_code == 303
    assert runner.calls == [
        (
            "comments",
            {"video_id": "7680023068660346011", "mode": "incremental"},
        ),
        ("comments", {"video_id": "7680023068660346011", "mode": "full"}),
    ]


def test_comment_count_route_is_distinct_from_comment_export(tmp_path: Path) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    video_id = archive_videos(store, 1)[0]
    runner = FakeRunner(store)

    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        response = client.post(
            f"/videos/{video_id}/comment-count",
            data={"csrf_token": csrf(client), "page": "3", "page_size": "100"},
            follow_redirects=False,
        )

    assert response.status_code == 303
    assert response.headers["location"] == (
        "/?page=3&page_size=100#video-archive"
    )
    assert runner.calls == [("comment_count", {"video_id": video_id})]


def test_comment_batch_returns_to_the_current_video_page(tmp_path: Path) -> None:
    local_settings = settings(tmp_path)
    store = LocalStore(local_settings.database_path)
    video_ids = archive_videos(store, 2)
    runner = FakeRunner(store)

    with TestClient(
        create_local_app(local_settings, store=store, runner=runner)
    ) as client:
        response = client.post(
            "/comments/batch",
            data={
                "csrf_token": csrf(client),
                "video_id": video_ids,
                "page": "3",
                "page_size": "100",
            },
            follow_redirects=False,
        )

    assert response.status_code == 303
    assert response.headers["location"] == (
        "/?page=3&page_size=100#video-archive"
    )
    assert runner.calls == [
        ("comments", {"video_id": video_id, "mode": "incremental"})
        for video_id in video_ids
    ]
