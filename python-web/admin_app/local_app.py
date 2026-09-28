"""FastAPI UI for the focused, file-first CSI OpenBase workflow."""

from __future__ import annotations

import asyncio
import hmac
import json
from contextlib import asynccontextmanager
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urlencode

from fastapi import FastAPI, Form, HTTPException, Request
from fastapi.responses import (
    FileResponse,
    HTMLResponse,
    JSONResponse,
    RedirectResponse,
)
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.base import BaseHTTPMiddleware
from starlette.middleware.sessions import SessionMiddleware
from starlette.middleware.trustedhost import TrustedHostMiddleware
from starlette.responses import Response

from . import __version__
from .local_cleanup import (
    CLEAR_DATA_LABELS,
    clear_local_data,
)
from .local_accounts import (
    CurrentRuntimeProxy,
    LocalAccountError,
    LocalAccountManager,
)
from .local_config import (
    COMMENT_EXPORT_DIRECTORY_KEY,
    LOCAL_PREFERENCES_META_KEY,
    LocalSettings,
    load_local_settings,
    prepare_comment_export_directory,
)
from .local_store import (
    GROUP_ID_RE,
    VIDEO_ID_RE,
    ActiveCommentJobError,
    ActiveLocalJobsError,
    LocalStore,
)
from .security import add_flash, csrf_token, pop_flashes, validate_csrf
from .time_utils import beijing_display, beijing_slug
from .viewmodels import pagination


APP_DIR = Path(__file__).resolve().parent
VIDEO_PAGE_SIZES = (30, 50, 100)
DEFAULT_VIDEO_PAGE_SIZE = VIDEO_PAGE_SIZES[0]
MAX_VIDEO_PAGE = 1_000_000
VIDEO_ARCHIVE_EXPORT_COLUMNS = (
    {"key": "video_id", "label": "视频 ID", "required": True},
    {"key": "title", "label": "标题", "required": True},
    {"key": "published_at", "label": "发布时间", "required": True},
    {"key": "description", "label": "描述", "required": False},
    {"key": "video_url", "label": "视频链接", "required": False},
    {"key": "groups", "label": "分组", "required": False},
    {"key": "view_count", "label": "播放数", "required": False},
    {"key": "like_count", "label": "点赞数", "required": False},
    {"key": "comment_count", "label": "评论数", "required": False},
    {"key": "collect_count", "label": "收藏数", "required": False},
    {"key": "share_count", "label": "分享数", "required": False},
    {"key": "first_seen_at", "label": "首次发现时间", "required": False},
    {"key": "last_seen_at", "label": "最近同步时间", "required": False},
    {"key": "observed_at", "label": "观测时间", "required": False},
    {"key": "sources", "label": "采集来源", "required": False},
)


def _with_time_displays(
    item: dict[str, Any], fields: tuple[str, ...]
) -> dict[str, Any]:
    value = dict(item)
    for field in fields:
        value[f"{field}_display"] = beijing_display(value.get(field))
    return value


def _video_latest_metadata(
    settings: LocalSettings, video: dict[str, Any]
) -> dict[str, Any]:
    record = video.get("record")
    indexed = dict(record) if isinstance(record, dict) else {}

    video_id = str(video.get("video_id") or "")
    if not VIDEO_ID_RE.fullmatch(video_id):
        return indexed
    video_directory = (settings.videos_dir / "douyin" / video_id).resolve()
    manifest_path = video_directory / "manifest.json"
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
        relative_metadata = Path(str(manifest.get("latest_metadata") or ""))
        metadata_path = (video_directory / relative_metadata).resolve()
        metadata_path.relative_to(video_directory)
        metadata = json.loads(metadata_path.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError, TypeError, json.JSONDecodeError):
        return indexed
    if not isinstance(metadata, dict):
        return indexed
    return {**indexed, **metadata}


def _video_published_at(settings: LocalSettings, video: dict[str, Any]) -> str | None:
    published_at = str(
        _video_latest_metadata(settings, video).get("published_at") or ""
    ).strip()
    return published_at or None


def _video_cover_path(
    settings: LocalSettings, video: dict[str, Any]
) -> Path | None:
    video_id = str(video.get("video_id") or "")
    relative_cover = str(video.get("cover_path") or "").strip()
    if not VIDEO_ID_RE.fullmatch(video_id) or not relative_cover:
        return None
    video_directory = (settings.videos_dir / "douyin" / video_id).resolve()
    candidate = (settings.data_home / relative_cover).resolve()
    try:
        candidate.relative_to(video_directory)
    except ValueError:
        return None
    return candidate if candidate.is_file() else None


def _video_archive_row(
    settings: LocalSettings, video: dict[str, Any]
) -> dict[str, Any]:
    metadata = _video_latest_metadata(settings, video)
    raw_metrics = metadata.get("visible_metrics")
    metrics = raw_metrics if isinstance(raw_metrics, dict) else {}

    def metric(name: str) -> int | None:
        value = metrics.get(name)
        return value if isinstance(value, int) and not isinstance(value, bool) else None

    source_labels = {"response": "平台接口", "dom": "页面可见内容"}
    raw_sources = metadata.get("sources")
    sources = raw_sources if isinstance(raw_sources, list) else []
    value = {
        **video,
        "description": str(metadata.get("desc") or "").strip(),
        "published_at": str(metadata.get("published_at") or "").strip() or None,
        "observed_at": str(metadata.get("observed_at") or "").strip() or None,
        "metrics": {
            "view_count": metric("view_count"),
            "like_count": metric("like_count"),
            "comment_count": metric("comment_count"),
            "collect_count": metric("collect_count"),
            "share_count": metric("share_count"),
        },
        "sources_display": " / ".join(
            source_labels.get(str(source), str(source))
            for source in sources
            if str(source).strip()
        )
        or "—",
        "cover_available": _video_cover_path(settings, video) is not None,
    }
    return _with_time_displays(
        value,
        ("published_at", "observed_at", "first_seen_at", "last_seen_at"),
    )


def _video_archive_export_value(video: dict[str, Any], key: str) -> Any:
    if key == "groups":
        return " / ".join(
            str(group.get("name") or "").strip()
            for group in video.get("groups") or ()
            if str(group.get("name") or "").strip()
        )
    if key in {"view_count", "like_count", "comment_count", "collect_count", "share_count"}:
        return dict(video.get("metrics") or {}).get(key)
    if key == "sources":
        value = str(video.get("sources_display") or "")
        return "" if value == "—" else value
    if key in {"published_at", "first_seen_at", "last_seen_at", "observed_at"}:
        return beijing_display(video.get(key), fallback="")
    return video.get(key, "")


def _video_archive_workbook(
    videos: list[dict[str, Any]], selected_columns: tuple[str, ...]
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill

    labels = {
        str(column["key"]): str(column["label"])
        for column in VIDEO_ARCHIVE_EXPORT_COLUMNS
    }
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("视频档案")
    sheet.freeze_panes = "A2"

    header = []
    for key in selected_columns:
        cell = WriteOnlyCell(sheet, value=labels[key])
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="24323B")
        header.append(cell)
    sheet.append(header)

    for video in videos:
        row = []
        for key in selected_columns:
            value = _video_archive_export_value(video, key)
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            row.append(cell)
        sheet.append(row)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _video_page(value: Any) -> int:
    try:
        parsed = int(str(value or "1"))
    except (TypeError, ValueError):
        return 1
    return min(max(parsed, 1), MAX_VIDEO_PAGE)


def _video_page_size(value: Any) -> int:
    try:
        parsed = int(str(value or DEFAULT_VIDEO_PAGE_SIZE))
    except (TypeError, ValueError):
        return DEFAULT_VIDEO_PAGE_SIZE
    return parsed if parsed in VIDEO_PAGE_SIZES else DEFAULT_VIDEO_PAGE_SIZE


def _video_return_path(form: Any) -> str:
    page = _video_page(form.get("page"))
    page_size = _video_page_size(form.get("page_size"))
    query: dict[str, Any] = {"page": page, "page_size": page_size}
    group_id = _video_group_filter(form.get("group_id"))
    if group_id:
        query["group_id"] = group_id
    return f"/?{urlencode(query)}#video-archive"


def _video_group_filter(value: Any) -> str | None:
    group_id = str(value or "").strip()
    if group_id == "ungrouped" or GROUP_ID_RE.fullmatch(group_id):
        return group_id
    return None


class DesktopTokenMiddleware(BaseHTTPMiddleware):
    """Require the per-launch desktop secret when one was configured."""

    def __init__(self, app: Any, *, token: str) -> None:
        super().__init__(app)
        self.token = token

    async def dispatch(self, request: Request, call_next: Any) -> Response:
        if not self.token or request.url.path == "/favicon.ico":
            return await call_next(request)
        supplied = request.headers.get("X-CSI-Desktop-Token", "")
        cookie = request.cookies.get("csi_desktop", "")
        if not (
            (supplied and hmac.compare_digest(supplied, self.token))
            or (cookie and hmac.compare_digest(cookie, self.token))
        ):
            return HTMLResponse("Desktop session authentication failed", status_code=401)
        response = await call_next(request)
        if supplied and not hmac.compare_digest(cookie, self.token):
            response.set_cookie(
                "csi_desktop",
                self.token,
                httponly=True,
                samesite="strict",
                secure=False,
            )
        return response


class SecurityHeadersMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request: Request, call_next: Any) -> Response:
        response = await call_next(request)
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["X-Frame-Options"] = "DENY"
        response.headers["Referrer-Policy"] = "no-referrer"
        response.headers["Content-Security-Policy"] = (
            "default-src 'self'; script-src 'self'; style-src 'self'; "
            "img-src 'self' data:; form-action 'self'; frame-ancestors 'none'"
        )
        response.headers["Cache-Control"] = "no-store"
        return response


class LocalRequestSerializationMiddleware(BaseHTTPMiddleware):
    """Keep account switching atomic relative to all local UI requests."""

    def __init__(self, app: Any) -> None:
        super().__init__(app)
        self._lock = asyncio.Lock()

    async def dispatch(self, request: Request, call_next: Any) -> Response:
        async with self._lock:
            return await call_next(request)


def _redirect(path: str = "/") -> RedirectResponse:
    return RedirectResponse(path, status_code=303)


def create_local_app(
    settings: LocalSettings | None = None,
    *,
    store: LocalStore | None = None,
    runner: Any | None = None,
    shutdown_callback: Any | None = None,
) -> FastAPI:
    root_settings = settings or load_local_settings()
    account_manager = LocalAccountManager(
        root_settings,
        store=store,
        runner=runner,
    )
    settings = CurrentRuntimeProxy(account_manager, "settings")
    store = CurrentRuntimeProxy(account_manager, "store")
    runner = CurrentRuntimeProxy(account_manager, "runner")

    @asynccontextmanager
    async def lifespan(_: FastAPI):
        account_manager.start()
        try:
            yield
        finally:
            account_manager.close()

    app = FastAPI(
        title="CSI OpenBase",
        version=__version__,
        docs_url=None,
        redoc_url=None,
        lifespan=lifespan,
    )
    app.state.local_settings = settings
    app.state.local_runner = runner
    app.state.local_accounts = account_manager
    app.add_middleware(LocalRequestSerializationMiddleware)
    app.add_middleware(SecurityHeadersMiddleware)
    app.add_middleware(
        DesktopTokenMiddleware, token=root_settings.desktop_token
    )
    app.add_middleware(
        SessionMiddleware,
        secret_key=root_settings.session_secret,
        same_site="strict",
        https_only=False,
    )
    app.add_middleware(
        TrustedHostMiddleware,
        allowed_hosts=["127.0.0.1", "localhost", "[::1]", "testserver"],
    )
    app.mount("/static", StaticFiles(directory=str(APP_DIR / "static")), name="static")
    templates = Jinja2Templates(directory=str(APP_DIR / "templates"))

    def account_context() -> dict[str, Any]:
        return {
            "account_workspaces": account_manager.account_summaries(),
            "active_workspace": account_manager.current_summary(),
        }

    def job_rows(limit: int = 30) -> list[dict[str, Any]]:
        return [
            _with_time_displays(job, ("created_at", "started_at", "finished_at"))
            for job in store.list_jobs(limit=limit)
        ]

    def page_context(request: Request) -> dict[str, Any]:
        jobs = job_rows(limit=1)
        requested_page = _video_page(request.query_params.get("page"))
        requested_page_size = _video_page_size(
            request.query_params.get("page_size")
        )
        requested_group_id = _video_group_filter(
            request.query_params.get("group_id")
        )
        groups = store.list_groups()
        known_groups = {str(group["group_id"]): group for group in groups}
        if requested_group_id not in {None, "ungrouped"} and requested_group_id not in known_groups:
            requested_group_id = None
        video_page = store.list_video_page(
            page=requested_page,
            page_size=requested_page_size,
            group_id=requested_group_id,
        )
        videos = [
            _with_time_displays(
                {
                    **video,
                    "published_at": _video_published_at(settings, video),
                },
                (
                    "published_at",
                    "first_seen_at",
                    "last_seen_at",
                    "last_comment_count_at",
                    "last_comment_export_at",
                ),
            )
            for video in video_page["items"]
        ]
        account = store.get_meta("creator_identity", {})
        last_discovery = store.get_meta("last_video_sync", {})
        last_export = _with_time_displays(
            store.get_meta("last_export", {}), ("finished_at",)
        )
        return {
            "request": request,
            "app_version": __version__,
            "csrf_token": csrf_token(request),
            "flashes": pop_flashes(request),
            "account": account,
            "authorized": bool(account and account.get("handle")),
            "jobs": jobs,
            "videos": videos,
            "video_total": video_page["total"],
            "video_page_size": video_page["page_size"],
            "groups": groups,
            "manual_groups": [group for group in groups if group["source"] == "manual"],
            "selected_group_id": requested_group_id,
            "selected_group": known_groups.get(str(requested_group_id or "")),
            "video_pagination": pagination(
                page=video_page["page"],
                total_pages=video_page["pages"],
                total_items=video_page["total"],
                path="/",
                query={
                    "page_size": video_page["page_size"],
                    **(
                        {"group_id": requested_group_id}
                        if requested_group_id
                        else {}
                    ),
                },
                fragment="video-archive",
            ),
            "active_jobs": store.active_job_count(),
            "data_home": str(settings.data_home),
            "last_export": last_export,
            "last_discovery": last_discovery,
            "video_synced": last_discovery.get("complete") is True,
            **account_context(),
        }

    @app.get("/health")
    def health() -> JSONResponse:
        return JSONResponse(
            {
                "status": "ok",
                "version": __version__,
                "mode": "local-archive",
                "active_jobs": store.active_job_count(),
                "account_id": account_manager.current.account.account_id,
                "account_count": len(
                    account_manager.registry.list(include_archived=False)
                ),
                "instance_nonce": settings.instance_nonce,
            }
        )

    @app.get("/", response_class=HTMLResponse)
    def home(request: Request) -> HTMLResponse:
        return templates.TemplateResponse(request, "local_home.html", page_context(request))

    @app.get("/video-archive", response_class=HTMLResponse)
    def video_archive(request: Request) -> HTMLResponse:
        requested_page = _video_page(request.query_params.get("page"))
        requested_page_size = _video_page_size(request.query_params.get("page_size"))
        video_page = store.list_video_page(
            page=requested_page,
            page_size=requested_page_size,
        )
        jobs = store.list_jobs(limit=1)
        account = store.get_meta("creator_identity", {})
        return templates.TemplateResponse(
            request,
            "local_video_archive.html",
            {
                "request": request,
                "app_version": __version__,
                "csrf_token": csrf_token(request),
                "account": account,
                "authorized": bool(account and account.get("handle")),
                "active_jobs": store.active_job_count(),
                "latest_job_id": jobs[0]["id"] if jobs else "",
                "videos": [
                    _video_archive_row(settings, video)
                    for video in video_page["items"]
                ],
                "video_total": video_page["total"],
                "video_page_size": video_page["page_size"],
                "video_export_columns": VIDEO_ARCHIVE_EXPORT_COLUMNS,
                "video_pagination": pagination(
                    page=video_page["page"],
                    total_pages=video_page["pages"],
                    total_items=video_page["total"],
                    path="/video-archive",
                    query={"page_size": video_page["page_size"]},
                ),
                **account_context(),
            },
        )

    @app.post("/video-archive/export")
    async def export_video_archive(request: Request) -> Response:
        form = await request.form()
        validate_csrf(request, form)
        known_columns = {
            str(column["key"]) for column in VIDEO_ARCHIVE_EXPORT_COLUMNS
        }
        requested_columns = {
            str(value).strip() for value in form.getlist("columns")
        }
        if not requested_columns.issubset(known_columns):
            raise HTTPException(status_code=400, detail="导出列无效")
        selected_columns = tuple(
            str(column["key"])
            for column in VIDEO_ARCHIVE_EXPORT_COLUMNS
            if column["required"] or column["key"] in requested_columns
        )
        videos = [
            _video_archive_row(settings, video)
            for video in store.list_videos(limit=100_000)
        ]
        body = _video_archive_workbook(videos, selected_columns)
        filename = f"csi-openbase-video-archive-{beijing_slug()}.xlsx"
        return Response(
            content=body,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Cache-Control": "no-store",
            },
        )

    @app.get("/video-covers/{video_id}", include_in_schema=False)
    def video_cover(video_id: str) -> FileResponse:
        if not VIDEO_ID_RE.fullmatch(video_id):
            raise HTTPException(status_code=404)
        video = store.get_video(video_id)
        if video is None:
            raise HTTPException(status_code=404)
        cover_path = _video_cover_path(settings, video)
        if cover_path is None:
            raise HTTPException(status_code=404)
        return FileResponse(cover_path)

    @app.get("/videos/{video_id}")
    def focus_video(request: Request, video_id: str) -> RedirectResponse:
        if not VIDEO_ID_RE.fullmatch(video_id):
            raise HTTPException(status_code=404)
        page = store.video_page_number(video_id, page_size=DEFAULT_VIDEO_PAGE_SIZE)
        if page is None:
            add_flash(request, "未找到对应的视频档案", "warning")
            return _redirect("/#video-archive")
        query = urlencode({"page": page, "page_size": DEFAULT_VIDEO_PAGE_SIZE})
        return _redirect(f"/?{query}#video-{video_id}")

    @app.get("/tasks", response_class=HTMLResponse)
    def task_history(request: Request) -> HTMLResponse:
        jobs = job_rows()
        account = store.get_meta("creator_identity", {})
        return templates.TemplateResponse(
            request,
            "local_tasks.html",
            {
                "request": request,
                "app_version": __version__,
                "account": account,
                "authorized": bool(account and account.get("handle")),
                "active_jobs": store.active_job_count(),
                "jobs": jobs,
                **account_context(),
            },
        )

    @app.get("/settings", response_class=HTMLResponse)
    def local_settings_page(request: Request) -> HTMLResponse:
        preferences = store.get_meta(LOCAL_PREFERENCES_META_KEY, {})
        jobs = store.list_jobs(limit=1)
        comment_export_directory = (
            preferences.get(COMMENT_EXPORT_DIRECTORY_KEY, "")
            if isinstance(preferences, dict)
            else ""
        )
        account = store.get_meta("creator_identity", {})
        return templates.TemplateResponse(
            request,
            "local_settings.html",
            {
                "request": request,
                "app_version": __version__,
                "csrf_token": csrf_token(request),
                "flashes": pop_flashes(request),
                "account": account,
                "authorized": bool(account and account.get("handle")),
                "active_jobs": store.active_job_count(),
                "latest_job_id": jobs[0]["id"] if jobs else "",
                "comment_export_directory": str(comment_export_directory or ""),
                "default_comment_directory": str(
                    settings.works_dir / "videos" / "douyin" / "<视频ID>" / "comments"
                ),
                **account_context(),
            },
        )

    @app.get("/accounts", response_class=HTMLResponse)
    def local_accounts_page(request: Request) -> HTMLResponse:
        account = store.get_meta("creator_identity", {})
        jobs = store.list_jobs(limit=1)
        return templates.TemplateResponse(
            request,
            "local_accounts.html",
            {
                "request": request,
                "app_version": __version__,
                "csrf_token": csrf_token(request),
                "flashes": pop_flashes(request),
                "account": account,
                "authorized": bool(account and account.get("handle")),
                "active_jobs": store.active_job_count(),
                "latest_job_id": jobs[0]["id"] if jobs else "",
                **account_context(),
            },
        )

    @app.post("/accounts")
    async def create_account(request: Request) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        if not account_manager.managed:
            raise HTTPException(status_code=409, detail="账号管理在测试模式下不可用")
        try:
            runtime = account_manager.create_and_switch(str(form.get("name") or ""))
        except (KeyError, LocalAccountError, OSError, RuntimeError, ValueError) as exc:
            add_flash(request, str(exc), "error")
            return _redirect("/accounts")
        add_flash(
            request,
            f"已创建并切换到“{runtime.account.name}”，请连接对应创作者账号",
            "success",
        )
        return _redirect("/")

    @app.post("/accounts/{account_id}/switch")
    async def switch_account(
        request: Request, account_id: str
    ) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        try:
            runtime = account_manager.switch(account_id)
        except (KeyError, LocalAccountError, OSError, RuntimeError, ValueError) as exc:
            add_flash(request, str(exc), "error")
            return _redirect("/accounts")
        add_flash(request, f"已切换到“{runtime.account.name}”", "success")
        return _redirect("/")

    @app.post("/accounts/{account_id}/rename")
    async def rename_account(
        request: Request, account_id: str
    ) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        try:
            account = account_manager.rename(account_id, str(form.get("name") or ""))
        except (KeyError, LocalAccountError, OSError, RuntimeError, ValueError) as exc:
            add_flash(request, str(exc), "error")
        else:
            add_flash(request, f"账号已重命名为“{account.name}”", "success")
        return _redirect("/accounts")

    @app.post("/accounts/{account_id}/archive")
    async def archive_account(
        request: Request, account_id: str
    ) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        try:
            account = account_manager.archive(account_id)
        except (KeyError, LocalAccountError, OSError, RuntimeError, ValueError) as exc:
            add_flash(request, str(exc), "error")
        else:
            add_flash(
                request,
                f"已从账号列表移除“{account.name}”，本地数据仍完整保留",
                "success",
            )
        return _redirect("/accounts")

    @app.post("/accounts/{account_id}/restore")
    async def restore_account(
        request: Request, account_id: str
    ) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        try:
            account = account_manager.restore(account_id)
        except (KeyError, LocalAccountError, OSError, RuntimeError, ValueError) as exc:
            add_flash(request, str(exc), "error")
        else:
            add_flash(request, f"已恢复“{account.name}”", "success")
        return _redirect("/accounts")

    @app.post("/settings/comments")
    async def save_comment_settings(request: Request) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        if store.active_job_count():
            add_flash(request, "任务运行时不可修改设置", "warning")
            return _redirect("/settings")
        try:
            selected = prepare_comment_export_directory(
                str(form.get("comment_export_directory") or ""), settings
            )
        except (OSError, ValueError) as exc:
            add_flash(request, str(exc), "error")
            return _redirect("/settings")

        preferences = store.get_meta(LOCAL_PREFERENCES_META_KEY, {})
        if not isinstance(preferences, dict):
            preferences = {}
        preferences[COMMENT_EXPORT_DIRECTORY_KEY] = (
            str(selected) if selected is not None else ""
        )
        store.set_meta(LOCAL_PREFERENCES_META_KEY, preferences)
        message = (
            f"评论将额外导出到 {selected}"
            if selected is not None
            else "已恢复默认，仅保存在工作目录的视频档案中"
        )
        add_flash(request, message, "success")
        return _redirect("/settings")

    def submit(
        request: Request,
        kind: str,
        *,
        redirect_path: str = "/",
        **payload: Any,
    ) -> RedirectResponse:
        try:
            job = runner.submit(kind, **payload)
        except ActiveCommentJobError as exc:
            add_flash(request, str(exc), "warning")
        except (KeyError, ValueError, RuntimeError) as exc:
            add_flash(request, str(exc), "error")
        else:
            add_flash(request, f"任务 #{job['id']} 已开始", "success")
        return _redirect(redirect_path)

    @app.post("/actions/authorize")
    async def authorize(request: Request) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        return submit(request, "authorize")

    @app.post("/actions/export")
    async def export_creator_data(request: Request) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        if not store.get_meta("creator_identity", {}):
            add_flash(request, "请先完成创作者账号授权", "error")
            return _redirect()
        last_video_sync = store.get_meta("last_video_sync", {})
        if not last_video_sync:
            add_flash(request, "请先同步作品档案，再导出全部表格", "error")
            return _redirect()
        if last_video_sync.get("complete") is not True:
            add_flash(request, "作品档案不完整，请重新同步后再导出", "error")
            return _redirect()
        return submit(request, "export")

    @app.post("/actions/sync-videos")
    async def sync_videos(request: Request) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        if not store.get_meta("creator_identity", {}):
            add_flash(request, "请先完成创作者账号授权", "error")
            return _redirect()
        return submit(request, "sync_videos")

    @app.post("/actions/clear-data")
    async def clear_data(request: Request) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        scope = str(form.get("scope") or "").strip()
        if scope not in CLEAR_DATA_LABELS:
            add_flash(request, "请选择有效的数据清理范围", "error")
            return _redirect()
        if form.get("confirm_clear") != "yes":
            add_flash(request, "请先确认清空操作无法撤销", "warning")
            return _redirect()
        try:
            result = clear_local_data(settings, store, scope)
        except ActiveLocalJobsError as exc:
            add_flash(request, str(exc), "warning")
        except (OSError, RuntimeError, ValueError) as exc:
            add_flash(request, f"清空失败：{exc}", "error")
        else:
            if result.pending_directories:
                add_flash(
                    request,
                    (
                        f"已从当前归档清空{CLEAR_DATA_LABELS[scope]}；"
                        f"仍有 {result.pending_directories} 个暂存目录因文件占用"
                        "未能删除，请关闭占用程序后再次清空"
                    ),
                    "warning",
                )
            else:
                add_flash(
                    request,
                    (
                        f"已清空{CLEAR_DATA_LABELS[scope]}："
                        f"删除 {result.files_deleted} 个文件、"
                        f"{result.directories_deleted} 个目录"
                    ),
                    "success",
                )
        return _redirect()

    @app.post("/groups")
    async def create_group(request: Request) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        try:
            group = store.create_group(str(form.get("name") or ""))
        except ValueError as exc:
            add_flash(request, str(exc), "error")
            return _redirect("/#video-archive")
        add_flash(request, f"已创建分组“{group['name']}”", "success")
        return _redirect(
            f"/?{urlencode({'group_id': group['group_id']})}#video-archive"
        )

    @app.post("/groups/{group_id}/rename")
    async def rename_group(request: Request, group_id: str) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        try:
            group = store.rename_group(group_id, str(form.get("name") or ""))
        except ValueError as exc:
            add_flash(request, str(exc), "error")
        else:
            add_flash(request, f"分组已重命名为“{group['name']}”", "success")
        return _redirect(_video_return_path(form))

    @app.post("/groups/{group_id}/delete")
    async def delete_group(request: Request, group_id: str) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        group = store.get_group(group_id)
        try:
            store.delete_group(group_id)
        except ValueError as exc:
            add_flash(request, str(exc), "error")
        else:
            name = str(group["name"]) if group else "该分组"
            add_flash(request, f"已删除分组“{name}”，作品档案仍然保留", "success")
        return _redirect("/#video-archive")

    async def change_group_videos(
        request: Request, *, add: bool
    ) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        raw_group_id = (
            form.get("target_group_id") if add else form.get("group_id")
        )
        group_id = str(raw_group_id or "").strip()
        video_ids = list(
            dict.fromkeys(str(value).strip() for value in form.getlist("video_id"))
        )
        try:
            if add:
                changed = store.add_videos_to_group(group_id, video_ids)
            else:
                changed = store.remove_videos_from_group(group_id, video_ids)
        except ValueError as exc:
            add_flash(request, str(exc), "error")
        else:
            action = "加入" if add else "移出"
            add_flash(request, f"已将 {changed} 个作品{action}分组", "success")
        return _redirect(_video_return_path(form))

    @app.post("/groups/videos/add")
    async def add_group_videos(request: Request) -> RedirectResponse:
        return await change_group_videos(request, add=True)

    @app.post("/groups/videos/remove")
    async def remove_group_videos(request: Request) -> RedirectResponse:
        return await change_group_videos(request, add=False)

    @app.post("/videos/{video_id}/comments")
    async def export_comments(request: Request, video_id: str) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        return submit(
            request,
            "comments",
            redirect_path=_video_return_path(form),
            video_id=video_id,
            mode="incremental",
        )

    @app.post("/videos/{video_id}/comments/full")
    async def synchronize_all_comments(
        request: Request, video_id: str
    ) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        return submit(
            request,
            "comments",
            redirect_path=_video_return_path(form),
            video_id=video_id,
            mode="full",
        )

    @app.post("/videos/{video_id}/comment-count")
    async def refresh_comment_count(
        request: Request, video_id: str
    ) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        return submit(
            request,
            "comment_count",
            redirect_path=_video_return_path(form),
            video_id=video_id,
        )

    @app.post("/comments/batch")
    async def export_comment_batch(request: Request) -> RedirectResponse:
        form = await request.form()
        validate_csrf(request, form)
        mode = str(form.get("mode") or "incremental").strip()
        if mode not in {"incremental", "full"}:
            add_flash(request, "评论导出模式无效", "error")
            return _redirect(_video_return_path(form))
        video_ids = list(
            dict.fromkeys(
                str(value).strip() for value in form.getlist("video_id")
            )
        )
        video_ids = [value for value in video_ids if value]
        if not video_ids:
            add_flash(request, "请选择至少一个视频", "error")
            return _redirect(_video_return_path(form))
        accepted = 0
        rejected = 0
        for video_id in video_ids:
            try:
                runner.submit("comments", video_id=video_id, mode=mode)
                accepted += 1
            except (ActiveCommentJobError, KeyError, ValueError, RuntimeError):
                rejected += 1
        level = "warning" if rejected else "success"
        action = "完整同步" if mode == "full" else "增量导出"
        add_flash(
            request,
            f"已启动 {accepted} 个评论{action}任务，跳过 {rejected} 个",
            level,
        )
        return _redirect(_video_return_path(form))

    @app.get("/api/state")
    def state() -> JSONResponse:
        jobs = store.list_jobs(limit=30)
        return JSONResponse(
            {
                "active_jobs": store.active_job_count(),
                "latest_job_id": jobs[0]["id"] if jobs else None,
                "jobs": jobs,
            }
        )

    @app.post("/api/shutdown")
    def shutdown() -> JSONResponse:
        if not root_settings.desktop_token:
            raise HTTPException(status_code=404)
        if shutdown_callback is None:
            raise HTTPException(status_code=503)
        force_required = account_manager.close() is False
        if not force_required:
            shutdown_callback()
        return JSONResponse(
            {"status": "stopping", "force_required": force_required}
        )

    @app.get("/favicon.ico", include_in_schema=False)
    def favicon() -> FileResponse:
        return FileResponse(
            APP_DIR / "static" / "brand" / "favicon.ico",
            media_type="image/x-icon",
        )

    return app
